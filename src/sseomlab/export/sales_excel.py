"""영업 키트 엑셀 출력 — 긴 문구는 줄바꿈 표시. 우선순위 정렬."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from sseomlab.sales.models import SalesKit

HEADERS = [
    "우선순위", "등급", "공간명", "지역", "우선연락시점", "한줄평가", "적합이유",
    "핵심이익", "평일유휴제안", "샘플촬영제안", "DM문구", "전화스크립트",
    "제안서제목", "추천상품명", "예상판매상품", "예상객단가",
]


def _row(k: SalesKit) -> list:
    return [
        k.contact_priority, k.grade, k.space_name, k.region, k.contact_timing,
        k.one_line_eval, k.why_fit, k.owner_benefit, k.weekday_idle_offer,
        k.sample_shoot_offer, k.dm_message, k.phone_script, k.proposal_title,
        k.product_name, " | ".join(k.expected_products), k.expected_avg_price_krw,
    ]


def export(kits: list[SalesKit], out_path: str | Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    kits = sorted(kits, key=lambda k: k.contact_priority)

    wb = Workbook()
    ws = wb.active
    ws.title = "영업키트"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="404040")
        c.alignment = Alignment(horizontal="center", vertical="center")

    wrap = Alignment(wrap_text=True, vertical="top")
    for k in kits:
        ws.append(_row(k))
        for c in ws[ws.max_row]:
            c.alignment = wrap

    ws.freeze_panes = "C2"
    widths = [9, 7, 16, 14, 22, 30, 38, 40, 32, 38, 46, 50, 30, 22, 26, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    wb.save(out_path)
    return out_path
