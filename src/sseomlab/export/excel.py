"""최종 결과를 엑셀(.xlsx)로 출력. 연락 우선순위 오름차순 정렬, 등급별 시각 강조."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from sseomlab.models import PlaceRecord

HEADERS = [
    "우선순위", "등급", "상호", "공간유형", "지역", "세부지역", "주소", "연락처",
    "네이버리뷰", "종합점수", "숨은공간", "웨딩전환", "돌상전환", "협업가능성",
    "추천제안", "네이버URL", "홈페이지", "인스타",
]

GRADE_FILL = {
    "A": PatternFill("solid", fgColor="C6EFCE"),
    "B": PatternFill("solid", fgColor="FFEB9C"),
    "C": PatternFill("solid", fgColor="FFF2CC"),
    "D": PatternFill("solid", fgColor="F2F2F2"),
}


def _row(r: PlaceRecord) -> list:
    p, a, s = r.place, r.analysis, r.score
    return [
        r.contact_priority, s.grade.value, p.name, a.space_type.value, p.region, p.sub_area,
        p.road_address or p.address, p.phone, p.naver_review_count, s.composite_score,
        s.hidden_space_score, s.wedding_conversion_score, s.dolsang_conversion_score,
        s.collaboration_probability, s.recommended_offer, p.naver_place_url, p.homepage, p.instagram_url,
    ]


def export(records: list[PlaceRecord], out_path: str | Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    records = sorted(records, key=lambda r: r.contact_priority)

    wb = Workbook()
    ws = wb.active
    ws.title = "발굴리스트"

    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="404040")
        c.alignment = Alignment(horizontal="center")

    for r in records:
        ws.append(_row(r))
        fill = GRADE_FILL.get(r.score.grade.value)
        if fill:
            ws.cell(row=ws.max_row, column=2).fill = fill

    ws.freeze_panes = "A2"
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 45)

    wb.save(out_path)
    return out_path
