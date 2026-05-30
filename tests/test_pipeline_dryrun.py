from pathlib import Path

from openpyxl import load_workbook

from sseomlab.pipeline import run


def test_dry_run_produces_excel(tmp_path: Path):
    out = tmp_path / "result.xlsx"
    path = run(dry_run=True, out_path=out, proposal_path=tmp_path / "prop.xlsx")
    assert path.exists()

    wb = load_workbook(path)
    ws = wb.active
    # 샘플 CSV 15건 중 웨딩홀/하우스웨딩 2건 제외 → 13건 + 헤더
    assert ws.max_row == 14
    assert ws.cell(row=2, column=1).value == 1     # 우선순위 1행
    assert ws.cell(row=2, column=2).value in ("A", "B")  # 등급
